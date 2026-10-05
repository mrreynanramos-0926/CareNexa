import json
from openpyxl import load_workbook

path = r"c:\Users\FBS80068\Downloads\Oct-2026-2.xlsx"
out = r"c:\Users\FBS80068\Documents\SourceCode\PYTHON_PROJECT\CareNexa\backend\src\modules\pricing\clinic-catalog.json"
wb = load_workbook(path, data_only=False)
ws = wb["TARAAN  18"]

def money(value):
    if isinstance(value, str) and value.startswith("="):
        return None
    return round(float(value) + 1e-9, 2)

items = []
seen = set()

def add(name, kind, price, cash, commission):
    clean = " ".join(str(name).split())
    if not clean or price is None:
        return
    key = (kind, clean.lower())
    if key in seen:
        clean = f"{clean} (2)"
        key = (kind, clean.lower())
    seen.add(key)
    row = {
        "name": clean,
        "type": kind,
        "price": f"{price:.2f}",
        "commissionPercent": commission,
        "description": None,
    }
    if cash is not None:
        row["cashPrice"] = f"{cash:.2f}"
    if kind == "service":
        row["description"] = "October 2026 list price. Gift-certificate cash is half. Gift-certificate card is 55% of the list price."
    else:
        row["description"] = "October 2026 list price. Cash uses the sheet cash price. Card is 95% of the list price."
    items.append(row)

# Services. Rows 5-71 are the 10% group, 72-81 the 15% group, 82-84 add-ons.
for row in range(5, 85):
    name = ws.cell(row, 1).value
    raw = ws.cell(row, 2).value
    if not isinstance(name, str):
        continue
    if isinstance(raw, str) and raw.startswith("="):
        # Package rows such as =B64*6 or =1500*6
        if "*6" in raw.replace(" ", ""):
            base_cell = raw.split("*")[0].replace("=", "").strip()
            if base_cell.startswith("B"):
                base = money(ws.cell(int(base_cell[1:]), 2).value)
            else:
                base = money(base_cell)
            price = None if base is None else round(base * 6, 2)
        else:
            price = None
    else:
        price = money(raw) if raw is not None else None
    if 5 <= row <= 71:
        commission = 10
    elif 72 <= row <= 81:
        commission = 15
    else:
        commission = 0
    add(name, "service", price, None, commission)

for row in range(96, 144):
    name = ws.cell(row, 1).value
    raw_price = ws.cell(row, 2).value
    raw_cash = ws.cell(row, 3).value
    if not isinstance(name, str) or raw_price is None:
        continue
    price = money(raw_price)
    if isinstance(raw_cash, str) and raw_cash.startswith("=") and "*0.9" in raw_cash.replace(" ", ""):
        cash = None if price is None else round(price * 0.9, 2)
    else:
        cash = money(raw_cash) if raw_cash is not None else None
    add(name, "product", price, cash, 15)

with open(out, "w", encoding="utf-8") as handle:
    json.dump(items, handle, indent=2)
    handle.write("\n")
print(len(items), "items")
print("services", sum(1 for item in items if item["type"] == "service"))
print("medicines", sum(1 for item in items if item["type"] == "product"))
